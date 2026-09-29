#!/usr/bin/env python3

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
import os
import tempfile

ROOT = Path(__file__).resolve().parent.parent
SOURCES = ROOT / "sources"
TEMPLATES = ROOT / "templates"
BRAND = ROOT / "brand"
FIXED_ZIP_TIME = (2026, 9, 29, 0, 0, 0)
PDF_FONT = Path("/System/Library/Fonts/STHeiti Medium.ttc")
BLUE = "17365D"
TEAL = "1F7A8C"
PALE = "EAF3F7"

def build_requirements(output: Path):
    doc = Document()
    configure_document(doc, "明澜科技项目交付协同平台需求说明")
    title = doc.add_paragraph(style="Title")
    title.add_run("明澜科技项目交付协同平台需求说明")
    intro = doc.add_paragraph("本文面向业务负责人、项目经理和信息化团队，明确首期范围、业务规则、确定性指标、验收边界和待确认事项。所有数据均为合成样例。")
    intro.style = doc.styles["Normal"]
    sections = [
        ("1 项目目标与范围", ["统一客户需求、方案评审、工作量估算、报价、计划和交付成果。", "首期服务销售顾问、方案架构师、项目经理和商务人员，共 260 名注册用户。", "不建设通用 CRM、财务核算、合同电子签章或公共模板交易市场。"]),
        ("2 角色与权限", ["销售顾问可创建项目和上传客户材料，但不能查看内部成本。", "方案架构师维护需求、功能和方案章节，可锁定已确认对象。", "商务人员维护角色单价和税率，只能发布经确认的报价版本。", "客户访客仅可访问短期分享链接中的客户可见成果。"]),
        ("3 材料与模板", ["项目材料支持 DOCX、PDF、PPTX、XLSX、CSV、TXT 和图片。", "企业模板支持 DOCX、XLSX 和 PPTX，系统只提取品牌颜色、字体和兼容页面尺寸。", "模板正文、宏、外部链接和嵌入对象不得复制到项目成果。"]),
        ("4 七类成果", ["系统形成需求分析、功能清单、工作量估算、实施计划、项目报价、整体解决方案和汇报演示。", "七类成果必须引用同一项目内容版本，并保持需求、功能、工作量、日期和金额一致。", "客户交付包不得包含内部成本、利润、原始材料或内部备注。"]),
        ("5 确定性计算", ["标准角色日单价为产品经理 1800 元、架构师 2200 元、后端工程师 1600 元、前端工程师 1500 元、测试工程师 1200 元。", "基准工作量为 324 人天，风险储备率 12%，税率 6%。", "未税实施金额等于各角色人天乘以日单价后求和，再乘以 1 加风险储备率；含税总额等于未税金额乘以 1 加税率。"]),
        ("6 计划与里程碑", ["计划从 2027 年 11 月 1 日开始，目标上线日期为 2028 年 5 月 31 日。", "里程碑依次为需求基线、方案评审、开发完成、试运行和正式上线。", "需求基线未确认时不得进入正式报价和上线承诺。"]),
        ("7 修改与版本", ["渲染变更只更新版式版本，不改变内容版本。", "内容变更必须生成影响计划，锁定对象不得被自动覆盖。", "回退必须恢复项目模型和受影响成果，并保留审计记录。"]),
        ("8 安全与部署", ["生产使用私有 PostgreSQL 和私有 MinIO，文件不得进入公共 Bucket。", "下载链接有效期 10 分钟，跨用户访问返回 404。", "上传文件先进入隔离区，扫描通过后才能进入正式对象存储。"]),
        ("9 服务目标", ["正常容量排队 P95 不超过 5 秒，突发容量不超过 30 秒。", "正式生成任务失败后自动重试，连续失败必须保留原因并停止伪健康心跳。", "单个项目完整生成时间目标为 30 分钟以内。"]),
        ("10 验收与排除", ["验收要求七类成果和成果包均可打开、可编辑、可下载且来源可追溯。", "关键数字错误、跨项目泄漏、文件无法打开或客户包泄漏内部成本均为阻断项。", "首期不包含 AWS S3、人工代写交付、自动购买基础设施或无人确认的外部动作。"]),
    ]
    for index, (heading, bullets) in enumerate(sections):
        doc.add_heading(heading, level=1)
        for bullet in bullets:
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(bullet)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        headers = ["验收编号", "检查内容", "结果要求"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            shade(cell, BLUE)
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
        for row_index in range(1, 4):
            cells = table.add_row().cells
            cells[0].text = f"AC-{index + 1:02d}-{row_index:02d}"
            cells[1].text = bullets[(row_index - 1) % len(bullets)]
            cells[2].text = "自动检查并保留来源" if row_index < 3 else "失败时阻止发布"
            if row_index % 2 == 0:
                for cell in cells:
                    shade(cell, PALE)
            for cell in cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        doc.add_paragraph()
    doc.core_properties.author = "Ningyi synthetic benchmark"
    doc.core_properties.created = fixed_datetime()
    doc.core_properties.modified = fixed_datetime()
    doc.save(output)
    normalize_office(output)


def build_docx_template(output: Path):
    doc = Document()
    configure_document(doc, "明澜科技企业方案模板")
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    title = doc.add_paragraph(style="Title")
    title.add_run("企业方案标题")
    doc.add_paragraph("客户名称  项目名称  版本日期")
    doc.add_page_break()
    doc.add_heading("章节标题", level=1)
    doc.add_paragraph("这里是模板示例正文，不属于任何项目事实，成果生成时不得复制。")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ["示例字段", "示例值", "说明"]):
        cell.text = text
        shade(cell, TEAL)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
    for row in table.rows[1:]:
        row.cells[0].text = "TEMPLATE-ONLY"
        row.cells[1].text = "999999"
        row.cells[2].text = "禁止进入正式成果"
    doc.core_properties.author = "Ningyi synthetic benchmark"
    doc.core_properties.created = fixed_datetime()
    doc.core_properties.modified = fixed_datetime()
    doc.save(output)
    patch_docx_theme(output)
    normalize_office(output)


def build_reference_pdf(output: Path):
    c = new_canvas(output, "整体方案参考说明")
    pages = [
        ("项目背景", ["明澜科技目前通过共享盘、邮件和即时消息协作项目方案。", "同一项目的需求、估算、报价和演示经常使用不同版本。", "目标是建立统一项目模型和可追溯成果链路。"]),
        ("目标业务链路", ["材料进入统一知识层后形成需求、功能和来源关系。", "确定性程序负责编号、工作量、金额、日期和版本。", "系统生成七类成果，并按照客户与内部权限形成成果包。"]),
        ("关键冲突", ["需求文档要求私有 MinIO，历史演示材料仍写 AWS S3。", "业务期望 2028 年 5 月上线，旧表格保留了 2028 年 3 月日期。", "报价必须以角色单价表和 12% 风险储备为准。"]),
        ("验收重点", ["模板仅提供视觉样式，不提供项目事实。", "七类成果中的 324 人天、12% 风险储备和 6% 税率保持一致。", "客户包不包含内部成本、利润、原始材料或模板示例正文。"]),
    ]
    for index, (title, bullets) in enumerate(pages):
        if index:
            c.showPage()
        draw_pdf_page(c, title, bullets, index + 1)
    c.save()


def build_brand_pdf(output: Path):
    c = new_canvas(output, "明澜科技品牌规范")
    draw_pdf_page(c, "明澜科技品牌规范", ["主色为深蓝 #17365D，辅色为青色 #1F7A8C，浅色背景为 #EAF3F7。", "中文正文字体使用 Noto Sans CJK SC，标题使用同一字体加粗。", "DOCX 默认 A4 纵向，PPTX 使用 16:9，XLSX 打印使用横向 A4。", "品牌规范只控制颜色、字体、页面尺寸和标识，不提供项目正文。"], 1)
    c.showPage()
    draw_pdf_page(c, "模板安全边界", ["模板示例文字、数字和表格仅用于展示样式。", "不得复制宏、外部链接、嵌入文件或隐藏内容。", "模板解析失败时回退平台默认版式并记录原因。", "正式成果必须保留原生可编辑文字和基础图形。"], 2)
    c.save()


def configure_document(doc: Document, title: str):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = styles[style_name]
        style.font.name = "Noto Sans CJK SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Heading 2"].font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].paragraph_format.space_after = Pt(18)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.2


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def new_canvas(output: Path, title: str):
    pdfmetrics.registerFont(TTFont("BenchmarkCJK", str(PDF_FONT), subfontIndex=0))
    result = canvas.Canvas(str(output), pagesize=A4, invariant=1, pageCompression=1)
    result.setTitle(title)
    result.setAuthor("Ningyi synthetic benchmark")
    return result


def draw_pdf_page(c, title: str, bullets: list[str], number: int):
    width, height = A4
    c.setFillColor(HexColor("#17365D"))
    c.setFont("BenchmarkCJK", 24)
    c.drawString(54, height - 82, title)
    c.setFillColor(HexColor("#627D98"))
    c.setFont("BenchmarkCJK", 9)
    c.drawRightString(width - 54, height - 54, f"合成基准材料  {number:02d}")
    y = height - 138
    c.setFillColor(HexColor("#243B53"))
    c.setFont("BenchmarkCJK", 11)
    for bullet in bullets:
        lines = wrap_text(bullet, 31)
        c.drawString(70, y, "•")
        for line in lines:
            c.drawString(88, y, line)
            y -= 25
        y -= 13
    c.setStrokeColor(HexColor("#D9E2EC"))
    c.line(54, 48, width - 54, 48)


def wrap_text(text: str, limit: int):
    return [text[index:index + limit] for index in range(0, len(text), limit)]


def fixed_datetime():
    from datetime import datetime, timezone
    return datetime(2026, 9, 29, tzinfo=timezone.utc)


def patch_docx_theme(path: Path):
    with ZipFile(path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}
    theme_name = "word/theme/theme1.xml"
    theme = files[theme_name].decode("utf-8")
    import re
    theme = re.sub(r'(<a:accent1>[\s\S]*?<a:srgbClr\s+val=")[0-9A-Fa-f]{6}', r'\g<1>17365D', theme, count=1)
    theme = re.sub(r'(<a:accent2>[\s\S]*?<a:srgbClr\s+val=")[0-9A-Fa-f]{6}', r'\g<1>1F7A8C', theme, count=1)
    theme = re.sub(r'(<a:accent3>[\s\S]*?<a:srgbClr\s+val=")[0-9A-Fa-f]{6}', r'\g<1>EAF3F7', theme, count=1)
    theme = re.sub(r'(<a:(?:latin|ea|cs)\b[^>]*\btypeface=")[^"]*', r'\g<1>Noto Sans CJK SC', theme)
    files[theme_name] = theme.encode("utf-8")
    write_zip(path, files)


def normalize_office(path: Path):
    with ZipFile(path, "r") as source:
        files = {name: source.read(name) for name in source.namelist()}
    if "docProps/core.xml" in files:
        core = files["docProps/core.xml"].decode("utf-8")
        import re
        core = re.sub(r"<dcterms:created[^>]*>.*?</dcterms:created>", '<dcterms:created xsi:type="dcterms:W3CDTF">2026-09-29T00:00:00Z</dcterms:created>', core)
        core = re.sub(r"<dcterms:modified[^>]*>.*?</dcterms:modified>", '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-09-29T00:00:00Z</dcterms:modified>', core)
        files["docProps/core.xml"] = core.encode("utf-8")
    write_zip(path, files)


def write_zip(path: Path, files: dict[str, bytes]):
    fd, temporary = tempfile.mkstemp(suffix=path.suffix, dir=str(path.parent))
    os.close(fd)
    try:
        with ZipFile(temporary, "w", ZIP_DEFLATED, compresslevel=6) as target:
            for name in sorted(files):
                info = ZipInfo(name, FIXED_ZIP_TIME)
                info.compress_type = ZIP_DEFLATED
                info.external_attr = 0o600 << 16
                target.writestr(info, files[name])
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


if __name__ == "__main__":
    for directory in (SOURCES, TEMPLATES, BRAND):
        directory.mkdir(parents=True, exist_ok=True)
    build_requirements(SOURCES / "platform-requirements.docx")
    build_docx_template(TEMPLATES / "enterprise-template.docx")
    build_reference_pdf(SOURCES / "solution-reference.pdf")
    build_brand_pdf(BRAND / "brand-guidelines.pdf")
